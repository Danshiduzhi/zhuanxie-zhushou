# -*- coding: utf-8 -*-
"""Word 文档与字幕生成"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_font(run, name_cn="宋体", name_en="Times New Roman", size=None,
             bold=None, color=None):
    run.font.name = name_en
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name_cn)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def setup_doc(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(4)
    for name, size, cn in (("Title", 18, "黑体"), ("Heading 1", 15, "黑体"),
                           ("Heading 2", 13, "黑体"), ("Heading 3", 12, "黑体")):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from docx.oxml import OxmlElement
    run = footer_p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(instr); run._r.append(f2)


def add_meta_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        set_font(c0.paragraphs[0].add_run(str(k)), name_cn="黑体",
                 size=Pt(10.5), bold=True)
        set_font(c1.paragraphs[0].add_run(str(v)), size=Pt(10.5))


def add_md_inline(p, text, bold=False, size=None):
    for j, part in enumerate(text.split("**")):
        if part:
            set_font(p.add_run(part), size=size, bold=bold or (j % 2 == 1))


def md_to_docx(md_text: str, out_path: Path, base_dir: Path | None = None):
    """总结 Markdown -> Word; 支持 ![图注](相对路径) 插图"""
    doc = Document()
    setup_doc(doc)
    lines = md_text.splitlines()
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # 图片: ![图注](路径)
        m_img = re.match(r"^\s*!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if m_img:
            cap, src = m_img.group(1), m_img.group(2)
            p = Path(src) if Path(src).is_absolute() else (base_dir or Path(".")) / src
            if p.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p_img.add_run()
                run.add_picture(str(p), width=Cm(15))
                if cap:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_font(p_cap.add_run(cap), size=Pt(10), color=RGBColor(0x66, 0x66, 0x66))
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows, i = _collect_table(lines, i)
            if rows:
                table = doc.add_table(rows=len(rows),
                                      cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(len(row)):
                        cell = table.rows[ri].cells[ci]
                        add_md_inline(cell.paragraphs[0], row[ci],
                                      bold=(ri == 0), size=Pt(10.5))
                doc.add_paragraph()
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], 3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_md_inline(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+[.、)]\s*", line):
            p = doc.add_paragraph(style="List Number")
            add_md_inline(p, re.sub(r"^\s*\d+[.、)]\s*", "", line))
        else:
            p = doc.add_paragraph()
            add_md_inline(p, line)
        i += 1
    doc.save(str(out_path))
    return out_path


def _collect_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def format_duration(sec: float) -> str:
    h = int(sec // 3600)
    m = int(sec % 3600 // 60)
    s = int(sec % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def build_record_docx(stem: str, meta: dict, transcript_text: str | None,
                      segments: list, doubts: list, out_path: Path):
    """文字记录 Word: 元信息 + 带时间戳/说话人正文 + 存疑汇总"""
    doc = Document()
    setup_doc(doc)
    content = stem.split("_", 2)[-1]
    doc.add_heading(content, 0)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run(f"{stem} · 文字记录"), size=Pt(12))
    add_meta_table(doc, [
        ("视频文件", meta.get("video", "-")),
        ("录制时间", meta.get("mtime", "-")),
        ("时长", format_duration(meta.get("duration", 0))),
        ("转写模型", meta.get("model", "-")),
        ("说话人数", str(meta.get("speakers", 0)) if meta.get("speakers") else "未启用"),
        ("存疑处数", str(len(doubts))),
    ])
    doc.add_heading("正文转录", 1)
    if transcript_text:
        for line in transcript_text.splitlines():
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph()
            m = re.match(r"^\[([\d:]+)\](\s*\[[A-Z0-9]+\])?\s?(.*)$", line)
            if m:
                set_font(p.add_run(f"[{m.group(1)}]"), size=Pt(11),
                         color=RGBColor(0x88, 0x88, 0x88))
                if m.group(2):
                    set_font(p.add_run(m.group(2)), size=Pt(11),
                             bold=True, color=RGBColor(0x1F, 0x6F, 0xB2))
                set_font(p.add_run(" " + m.group(3)), size=Pt(11))
            else:
                set_font(p.add_run(line), size=Pt(11))
    else:
        set_font(doc.add_paragraph().add_run("(无有效语音内容)"), size=Pt(11))
    doc.add_heading("存疑与待确认", 1)
    if doubts:
        set_font(doc.add_paragraph().add_run(
            "以下片段语音置信度低或语义存疑, 建议对照原音频复核:"), size=Pt(11))
        for d in doubts:
            p = doc.add_paragraph()
            t = d.get("t", "")
            set_font(p.add_run(f"[{t}] " if t else ""), size=Pt(11),
                     color=RGBColor(0x88, 0x88, 0x88))
            set_font(p.add_run(d.get("text", "")), size=Pt(11))
    else:
        set_font(doc.add_paragraph().add_run("本段无存疑片段。"), size=Pt(11))
    doc.add_heading("备注", 1)
    set_font(doc.add_paragraph().add_run(
        "标记说明: [存疑] 表示置信度低或语义不明; [A]/[B] 为说话人标签; "
        "[MM:SS]/[HH:MM:SS] 为时间戳, 可在音频/视频中定位。"), size=Pt(10.5))
    doc.save(str(out_path))
    return out_path


def build_srt(segments: list, total: float, out_path: Path):
    """字幕文件 .srt"""
    lines = []
    for i, s in enumerate(segments, 1):
        st, en = s["start"], s["end"]
        hh1, mm1, ss1 = int(st // 3600), int(st % 3600 // 60), st % 60
        hh2, mm2, ss2 = int(en // 3600), int(en % 3600 // 60), en % 60
        tag = f"[{s['speaker']}] " if s.get("speaker") else ""
        lines.append(f"{i}\n{hh1:02d}:{mm1:02d}:{ss1:06.3f} --> "
                     f"{hh2:02d}:{mm2:02d}:{ss2:06.3f}\n{tag}{s['text']}\n")
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path
